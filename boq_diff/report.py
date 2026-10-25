"""Report writers for boq-diff: Markdown and colour-highlighted Excel."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from .core import Change, DiffResult, Item, total_amount


def _fmt(value: Decimal) -> str:
    return f"{value:,.2f}"


def _signed(value: Decimal) -> str:
    if value > 0:
        return f"+{_fmt(value)}"
    return _fmt(value)


def summary_lines(old_items, new_items, result: DiffResult) -> list[str]:
    old_total = total_amount(old_items)
    new_total = total_amount(new_items)
    delta = new_total - old_total
    return [
        f"- 旧版清单：{len(old_items)} 项，合价总计 {_fmt(old_total)} 元",
        f"- 新版清单：{len(new_items)} 项，合价总计 {_fmt(new_total)} 元",
        f"- 总价变化：{_signed(delta)} 元",
        f"- 新增 {len(result.added)} 项 / 删除 {len(result.removed)} 项 / "
        f"变更 {len(result.changed)} 项 / 未变 {result.unchanged_count} 项",
    ]


def _item_row(prefix: str, item: Item) -> str:
    return (
        f"| {prefix} | {item.code or '-'} | {item.name} | {item.unit or '-'} "
        f"| {_fmt(item.quantity)} | {_fmt(item.unit_price)} | {_fmt(item.amount)} |"
    )


def render_markdown(old_items, new_items, result: DiffResult,
                     old_label: str = "旧版", new_label: str = "新版") -> str:
    lines = [
        "# 工程量清单差异报告",
        "",
        f"对比：{old_label} → {new_label}",
        "",
        "## 汇总",
        "",
        *summary_lines(old_items, new_items, result),
        "",
    ]
    if not result.has_changes:
        lines += ["两版清单完全一致，没有任何差异。", ""]
        return "\n".join(lines)

    header = "| | 编码 | 项目名称 | 单位 | 工程量 | 综合单价 | 合价 |\n|---|---|---|---|---|---|---|"
    if result.added:
        lines += ["## 新增项", "", header]
        lines += [_item_row("🟢 新增", it) for it in result.added]
        lines.append("")
    if result.removed:
        lines += ["## 删除项", "", header]
        lines += [_item_row("🔴 删除", it) for it in result.removed]
        lines.append("")
    if result.changed:
        lines += [
            "## 变更项",
            "",
            "| 编码 | 项目名称 | 单位 | 工程量（旧 → 新 / 变化） | 综合单价（旧 → 新 / 变化） | 合价变化 |",
            "|---|---|---|---|---|---|",
        ]
        for ch in result.changed:
            o, n = ch.old, ch.new
            lines.append(
                f"| {n.code or '-'} | {n.name} | {n.unit or '-'} "
                f"| {_fmt(o.quantity)} → {_fmt(n.quantity)}（{_signed(ch.quantity_delta)}） "
                f"| {_fmt(o.unit_price)} → {_fmt(n.unit_price)}（{_signed(ch.unit_price_delta)}） "
                f"| {_signed(ch.amount_delta)} |"
            )
        lines.append("")
    return "\n".join(lines)


def write_xlsx(path, old_items, new_items, result: DiffResult) -> Path:
    """Write one Excel workbook: a 汇总 sheet plus a colour-coded 明细 sheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError(
            "生成 .xlsx 报告需要 openpyxl，请先安装: pip install openpyxl"
        ) from exc

    wb = Workbook()

    ws = wb.active
    ws.title = "汇总"
    ws.append(["工程量清单差异报告"])
    ws["A1"].font = Font(bold=True, size=14)
    for line in summary_lines(old_items, new_items, result):
        ws.append([line.lstrip("- ")])
    ws.column_dimensions["A"].width = 60

    detail = wb.create_sheet("差异明细")
    headers = [
        "类型", "编码", "项目名称", "单位",
        "旧工程量", "新工程量", "工程量变化",
        "旧综合单价", "新综合单价", "单价变化",
        "旧合价", "新合价", "合价变化",
    ]
    detail.append(headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)

    fills = {
        "新增": PatternFill("solid", fgColor="C6EFCE"),
        "删除": PatternFill("solid", fgColor="FFC7CE"),
        "变更": PatternFill("solid", fgColor="FFEB9C"),
    }

    def _row(kind: str, o: Item | None, n: Item | None) -> list:
        ref = n or o
        return [
            kind, ref.code, ref.name, ref.unit,
            float(o.quantity) if o else None,
            float(n.quantity) if n else None,
            float(n.quantity - o.quantity) if o and n else None,
            float(o.unit_price) if o else None,
            float(n.unit_price) if n else None,
            float(n.unit_price - o.unit_price) if o and n else None,
            float(o.amount) if o else None,
            float(n.amount) if n else None,
            float(n.amount - o.amount) if o and n else None,
        ]

    rows = (
        [("新增", None, it) for it in result.added]
        + [("删除", it, None) for it in result.removed]
        + [("变更", ch.old, ch.new) for ch in result.changed]
    )
    for kind, o, n in rows:
        detail.append(_row(kind, o, n))
        fill = fills[kind]
        for cell in detail[detail.max_row]:
            cell.fill = fill

    widths = [8, 16, 24, 8, 12, 12, 12, 12, 12, 12, 14, 14, 14]
    for idx, w in enumerate(widths, start=1):
        from openpyxl.utils import get_column_letter
        detail.column_dimensions[get_column_letter(idx)].width = w

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
